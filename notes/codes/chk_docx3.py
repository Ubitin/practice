# -*- coding: utf-8 -*-
# 检查打印版 docx 里 §6.16 是否被正确识别为标题
import glob, os, re, sys

d = r"D:\lenovo\chat_with_deepseek_harness\讲义"
cands = [f for f in glob.glob(os.path.join(d, "*.docx"))]
print("找到 docx:")
for c in cands:
    print("  ", os.path.basename(c), os.path.getsize(c))

try:
    from docx import Document
except Exception as e:
    print("python-docx 不可用:", e)
    sys.exit(1)

p = [c for c in cands if "打印版" in os.path.basename(c)][0]
doc = Document(p)
print("\n段落总数 =", len(doc.paragraphs))
hits = []
for i, para in enumerate(doc.paragraphs):
    t = para.text.strip()
    if "6.16" in t or "6.15" in t or "组合数" in t[:12] or "单调栈" in t[:12]:
        hits.append((i, para.style.name, t[:60]))
print("\n相关段落:")
for i, st, t in hits[:15]:
    print("  #%-5d [%-12s] %s" % (i, st, t))

# 统计各样式
from collections import Counter
c = Counter(pp.style.name for pp in doc.paragraphs)
print("\n样式分布:", dict(c))
