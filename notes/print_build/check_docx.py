#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""检查生成的 DOCX：纸张/边距/页脚域/标题字体/代码底纹/内容抽样。"""
import sys
from collections import Counter

from docx import Document
from docx.oxml.ns import qn

path = sys.argv[1]
d = Document(path)
s = d.sections[0]
print(f'page: {s.page_width.mm:.1f} x {s.page_height.mm:.1f} mm (A4=210x297)')
print(f'margins mm: T{s.top_margin.mm:.1f} B{s.bottom_margin.mm:.1f} '
      f'L{s.left_margin.mm:.1f} R{s.right_margin.mm:.1f}')
xml = s.footer.paragraphs[0]._p.xml
print('footer PAGE:', 'PAGE' in xml, '| NUMPAGES:', 'NUMPAGES' in xml)
print('updateFields:', d.settings.element.find(qn('w:updateFields')) is not None)

cnt = Counter()
for p in d.paragraphs:
    try:
        cnt[p.style.name] += 1
    except Exception:
        cnt['?'] += 1
print('paragraphs:', dict(cnt))

# 直接格式化检查：标题 run 字体、代码 run 字体
def first_run_font(p):
    for r in p.runs:
        rpr = r._element.find(qn('w:rPr'))
        if rpr is None:
            continue
        rf = rpr.find(qn('w:rFonts'))
        if rf is not None:
            ea = rf.get(qn('w:eastAsia')) or ''
            ascii_f = rf.get(qn('w:ascii')) or ''
            return ascii_f, ea, (r.font.size.pt if r.font.size else None)
    return None

hits = {}
for p in d.paragraphs:
    try:
        nm = p.style.name
    except Exception:
        continue
    if nm in ('Heading 1', 'Heading 2', 'Heading 3', 'Source Code') and nm not in hits:
        hits[nm] = (p.text[:24], first_run_font(p))
for k, v in hits.items():
    print(f'{k:12s} run={v[1]}  text={v[0]!r}')

# 底纹段数与内容抽样
shaded = sum(1 for p in d.paragraphs
             if p._p.find(qn('w:pPr')) is not None
             and p._p.find(qn('w:pPr')).find(qn('w:shd')) is not None)
print('shaded paragraphs:', shaded)
alltext = '\n'.join(p.text for p in d.paragraphs)
for probe in ('#include <bits/stdc++.h>', 'bitset<1000005>', 'memset(arr, 0, sizeof(arr))',
              '誊抄建议', '机试代码模板速查手册（打印版）'):
    print(f'  contains {probe!r}:', probe in alltext)
print('total chars:', len(alltext))
